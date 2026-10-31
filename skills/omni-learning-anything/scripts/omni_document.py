"""Shared content contract and standalone Markdown/DOCX export. PDF imports are lazy."""
from omni_locale import source_labels
import html
import re
import shutil
import subprocess
import sys
from pathlib import Path
from urllib.parse import quote, urlsplit
import omni_core as c
from omni_locale import language, subtitle, CJK_FONTS, ARABIC_FONTS


def validate_document(data,root):
    sources=data.get('sources',[])
    config=c.read(root/'config.json'); material_config=config.get('materials')
    if material_config:
        from omni_materials import validate_config as validate_material_config,validate_source
        validate_material_config(material_config)
        if not sources: raise ValueError('At least one actual evidence source is required.')
        bound=[s for s in sources if 'material_ref' in s]
        if material_config['policy']=='strict' and len(bound)!=len(sources):raise ValueError('Strict material mode cannot introduce unbound external sources.')
        if material_config['policy']=='supplemented' and not bound:raise ValueError('Material-first mode requires selected-material evidence.')
        for source in bound:validate_source(root,source,config)
    elif any('material_ref' in s for s in sources):raise ValueError('Material references require explicit configured material mode.')
    elif len({s.get('url') for s in sources if s.get('url')})<5: raise ValueError('At least five distinct expansion links required.')
    ids=[s.get('id') for s in sources]
    if len(set(ids))!=len(ids) or any(not re.fullmatch(r'S\d+',str(x)) for x in ids): raise ValueError('Source IDs must be unique S-numbers.')
    for s in sources:
        u=urlsplit(s.get('url',''))
        if material_config and not s.get('url') and s.get('material_ref'):pass
        elif u.scheme not in ('http','https') or not u.netloc or any(ch in s['url'] for ch in '\\{}\n\r'): raise ValueError('Invalid source URL.')
        if not all(s.get(k) for k in ('title','publisher','note','checked_at','published_at')): raise ValueError('Incomplete source metadata.')
        if s['checked_at']!=c.today(root): raise ValueError('Recheck sources on the actual generation date.')
    if not data.get('sections') or not all(s.get('title') and s.get('body') for s in data['sections']): raise ValueError('Sections need title and body.')
    text='\n'.join(s['body'] for s in data['sections'])
    if not set(re.findall(r'\[(S\d+)\]',text)).issubset(ids): raise ValueError('Unknown inline source.')
    if not re.search(r'\[S\d+\]',text): raise ValueError('Include inline evidence for key claims.')
    f=data.get('figure',{}); p=(root/f.get('path','')).resolve()
    if not p.is_relative_to(root.resolve()) or not p.is_file() or p.suffix.lower() not in ('.png','.jpg','.jpeg','.pdf'): raise ValueError('Figure must be a local supported file inside project.')
    if not f.get('caption') or not f.get('credit'): raise ValueError('Figure caption and credit required.')
    if not data.get('title'): raise ValueError('Document title required.')


def validate_key(key):
    if not re.fullmatch(r'plan|Day[0-9]{2,}', key) or (key != 'plan' and int(key[3:]) < 1):
        raise ValueError('Invalid document key.')


def prepare(root, input_path, key):
    validate_key(key)
    if not input_path.is_relative_to(root):
        raise ValueError('Input must be inside project.')
    state = c.read(root / 'state.json'); plan = c.read(root / 'plan.json')
    c.validate_plan(plan)
    config = c.read(root / 'config.json'); data = c.read(input_path)
    if key == 'plan':
        if input_path != root / 'plan.json':
            raise ValueError('Plan renderer must use canonical plan.json.')
        if state['lessons']:
            raise ValueError('Preserve the started curriculum; create a new project.')
        if (root / 'manifests/plan.json').exists():
            raise ValueError('Plan already rendered; reuse it or create a revised project.')
    else:
        c.check_approval(root, state)
        day = data.get('day'); action = c.next_action(root)
        if type(day) is not int or key != f'Day{day:02d}' or action.get('day') != day or action['action'] != 'generate':
            raise ValueError('Render only the next ungenerated lesson; reuse existing artifacts.')
        if data.get('title') != plan['days'][day - 1]['topic']:
            raise ValueError('Lesson title differs from approved plan.')
        reading = data.get('reading_minutes'); practice = data.get('practice_minutes')
        if type(reading) is not int or reading < 1 or type(practice) is not int or practice < 0:
            raise ValueError('Reading minutes must be positive; practice minutes nonnegative integers.')
        if reading + practice > config['daily_minutes']:
            raise ValueError('Lesson exceeds agreed time budget.')
    validate_document(data, root)
    return data, plan, config


def material_markdown(value):
    # A literal user excerpt must not introduce active HTML, images or Markdown syntax.
    return re.sub(r'([\\`*_{}\[\]()#+.!|>])', r'\\\1', html.escape(value, quote=False))


def markdown_text(data, plan, config, date, key, image_ref):
    labels=source_labels(config)
    lines = ['# ' + data['title'], subtitle(data, plan, config, date, key)]
    for i, section in enumerate(data['sections']):
        lines.extend(['## ' + section['title'], section['body']])
        if i == min(1, len(data['sections']) - 1):
            lines.extend([f"![{data['figure']['caption']}]({image_ref})", data['figure']['credit']])
    lines.append('## '+labels['further'])
    for source in data['sources']:
        title=material_markdown(source['title']) if source.get('material_ref') else source['title']
        note=material_markdown(source['note']) if source.get('material_ref') else source['note']
        label=f"[{title}]({source['url']})" if source.get('url') else title
        lines.append(f"- [{source['id']}] {label}: {note}; {source['publisher']} | {labels['published']}: {source['published_at']} | {labels['checked']}: {source['checked_at']}")
    return '\n\n'.join(lines) + '\n'


def cjk_font(match=None, language='zh'):
    candidates = CJK_FONTS.get(language,CJK_FONTS['zh'])
    fallback = ('Hiragino Mincho ProN' if sys.platform=='darwin' else 'Noto Serif CJK JP') if language=='ja' else (('Apple SD Gothic Neo' if sys.platform=='darwin' else 'Noto Serif CJK KR') if language=='ko' else ('STSong' if sys.platform=='darwin' else ('SimSun' if sys.platform=='win32' else 'Noto Serif CJK SC')))
    if match is None:
        binary = shutil.which('fc-match')
        if binary:
            def match(name):
                result = subprocess.run([binary, '-f', '%{family}', name], capture_output=True, text=True, timeout=5)
                return result.stdout if result.returncode == 0 else ''
        else:
            # Fontconfig is optional. Use platform-native CJK families when unavailable.
            return fallback
    for name in candidates:
        families = [s.strip().casefold() for s in match(name).split(',')]
        if name.casefold() in families:
            return name
    # DOCX stores Unicode/font declarations; cloud viewers may have fonts the authoring
    # host lacks. Rendering remains a separate inspection requirement.
    return fallback


def write_docx(path, data, plan, config, image, date, key):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from PIL import Image
    labels=source_labels(config)
    document = Document()
    document.core_properties.author = 'Omni Learning Anything'
    document.core_properties.title = data['title']
    east_asian_font = cjk_font(language=labels['code'])
    complex_font = 'Geeza Pro' if sys.platform=='darwin' else 'Noto Naskh Arabic'
    for style in document.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = document.styles[name]
        style.font.name = 'Times New Roman'; style.font.size = Pt(11 if name == 'Normal' else 16)
        fonts = style.element.get_or_add_rPr().rFonts
        fonts.set(qn('w:eastAsia'), east_asian_font)
        if labels['dir']=='rtl':fonts.set(qn('w:cs'),complex_font)
        lang=OxmlElement('w:lang');lang.set(qn('w:val'),labels['code']);lang.set(qn('w:bidi'),labels['code']);style.element.get_or_add_rPr().append(lang)
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + attr), None)
    document.styles['Normal'].paragraph_format.space_after = Pt(7)
    document.add_heading(data['title'], 0)
    document.add_paragraph(subtitle(data, plan, config, date, key))
    for i, item in enumerate(data['sections']):
        document.add_heading(item['title'], 1)
        for paragraph in item['body'].split('\n\n'):
            document.add_paragraph(paragraph)
        if i == min(1, len(data['sections']) - 1):
            with Image.open(image) as img:
                ratio = img.height / img.width
            width = min(6.5, 3.0 / ratio)
            shape = document.add_picture(str(image), width=Inches(width))
            shape._inline.docPr.set('descr', data['figure']['caption'])
            document.add_paragraph(data['figure']['caption'])
            document.add_paragraph(data['figure']['credit'])
    document.add_heading(labels['further'], 1)
    for source in data['sources']:
        p = document.add_paragraph(f"[{source['id']}] ")
        if not source.get('url'):
            p.add_run(source['title']+f": {source['note']}; {source['publisher']} | {labels['published']}: {source['published_at']} | {labels['checked']}: {source['checked_at']}")
            continue
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(source['url'], RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r'); properties = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '227E82'); properties.append(color)
        run.append(properties); text = OxmlElement('w:t'); text.text = source['title']; run.append(text)
        link.append(run); p._p.append(link)
        p.add_run(f": {source['note']}; {source['publisher']} | {labels['published']}: {source['published_at']} | {labels['checked']}: {source['checked_at']}")
    if labels['dir']=='rtl':
        from copy import deepcopy
        arabic=r'[\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff]'
        for paragraph in document.paragraphs:
            properties=paragraph._p.get_or_add_pPr()
            bidi=OxmlElement('w:bidi');bidi.set(qn('w:val'),'1');properties.append(bidi)
            # Logical start stays at the right edge in RTL-aware Word/LO viewers.
            jc=properties.find(qn('w:jc'))
            if jc is None:jc=OxmlElement('w:jc');properties.append(jc)
            jc.set(qn('w:val'),'start')
        for run in list(document.element.xpath('.//w:r')):
            nodes=run.findall(qn('w:t'))
            if len(nodes)!=1:continue
            text=nodes[0].text or '';parent=run.getparent();index=parent.index(run)
            parts=re.split('('+arabic+'+(?:[ \t]+'+arabic+'+)*)',text)
            for part in parts:
                if not part:continue
                new=deepcopy(run);new.find(qn('w:t')).text=re.sub(r'\b[0-9]{4}-[0-9]{2}-[0-9]{2}\b',lambda m:'\u200e'+m.group()+'\u200e',part);new.find(qn('w:t')).set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                rtl=OxmlElement('w:rtl');rtl.set(qn('w:val'),'1' if re.search(arabic,part) else '0');new.get_or_add_rPr().append(rtl)
                parent.insert(index,new);index+=1
            parent.remove(run)
    document.save(path)


def render(root, input_path, key, formats=None):
    root = Path(root).resolve(); input_path = Path(input_path).resolve()
    formats = list(dict.fromkeys(formats or ['markdown']))
    if not formats or any(f not in ('pdf', 'markdown', 'docx', 'html') for f in formats):
        raise ValueError('Choose pdf, markdown, docx or html formats.')
    validate_key(key)
    if 'pdf' in formats:
        import omni_pdf
        return omni_pdf.render(root, input_path, key, export_formats=formats)
    with c.lock(root):
        data, plan, config = prepare(root, input_path, key)
        for folder in ('sources', 'assets', 'manifests', 'previews'):
            path = root / folder
            if not path.resolve().is_relative_to(root):
                raise ValueError('Output directory escapes project.')
            path.mkdir(exist_ok=True)
        figure = (root / data['figure']['path']).resolve()
        if figure.suffix.lower() == '.pdf':
            raise ValueError('Markdown/DOCX figures need PNG or JPEG; convert the figure first.')
        image = root / 'assets' / ('figure-' + c.digest(figure)[:12] + figure.suffix.lower())
        if image != figure:
            shutil.copyfile(figure, image)
        artifacts = {}; outputs = [image]
        md = root / 'sources' / (key + '.md')
        text = markdown_text(data, plan, config, c.today(root), key, '../assets/' + quote(image.name))
        md.write_text(text, encoding='utf-8'); artifacts['markdown'] = str(md.relative_to(root)); outputs.append(md)
        if 'docx' in formats:
            path = root / 'sources' / (key + '.docx')
            write_docx(path, data, plan, config, image, c.today(root), key)
            artifacts['docx'] = str(path.relative_to(root)); outputs.append(path)
        preview = root / 'previews' / (key + '.html')
        from omni_html import serialize
        html_text=serialize(data,plan,config,image,c.today(root),key)
        preview.write_text(html_text,encoding='utf-8')
        if 'html' in formats:
            path=root/'sources'/(key+'.html');path.write_text(html_text,encoding='utf-8')
            artifacts['html']=str(path.relative_to(root));outputs.append(path)
        outputs.extend([preview, input_path, root / 'plan.json', root / 'config.json'])
        manifest = {'key': key, 'artifact': artifacts[formats[0]],
                    'artifacts': artifacts, 'input': str(input_path.relative_to(root)),
                    'created_at': c.today(root), 'files': {str(p.relative_to(root)): c.digest(p) for p in outputs},
                    'previews': [str(preview.relative_to(root))], 'reviewed_at': None,
                    'report': {'links': len(data['sources']), 'figure': str(image.relative_to(root))}}
        c.save(root / 'manifests' / (key + '.json'), manifest)
        return manifest


def review(root, key, note):
    root = Path(root).resolve(); validate_key(key)
    with c.lock(root):
        state = c.read(root / 'state.json')
        if key != 'plan' and str(int(key[3:])) in state['lessons']:
            raise ValueError('Delivered review is immutable; retain the original archive.')
        path = root / 'manifests' / (key + '.json'); manifest = c.read(path)
        c.manifest_ok(root, manifest)
        if not note.strip():
            raise ValueError('Record what was visually checked.')
        manifest.update(reviewed_at=c.today(root), review_note=note)
        c.save(path, manifest)
